# backend/report/views.py
from rest_framework.generics import ListAPIView
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from rest_framework.views import APIView
from rest_framework.pagination import PageNumberPagination
from django.db.models import Count, Q, Exists, OuterRef, Prefetch 
from loan.models import Loan, LoanRoleAssignment, Task, Milestone
from employee.models import Employee
from employee.utils import to_cst
import logging
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from django.http import HttpResponse
from openpyxl import Workbook
from django.shortcuts import get_object_or_404
from openpyxl.styles import Font
from datetime import datetime
from django.db.models.functions import ExtractMonth
from django.http import HttpResponse
from datetime import datetime
from openpyxl.utils import get_column_letter


logger = logging.getLogger(__name__)


# -----------------------------
# Pagination
# -----------------------------
class StandardResultsSetPagination(PageNumberPagination):
    page_size = 10
    page_size_query_param = "page_size"
    max_page_size = 10000


class FundedLoanReportAPIView(ListAPIView):
    permission_classes = [IsAuthenticated]
    pagination_class = None

    def get_queryset(self):
        user = self.request.user
        qs = Loan.objects.all().select_related("broker", "loan_officer", "milestone")

        # Visibility for non-admin users
        if not user.is_superuser and not user.has_perm("loan.view_all_loans"):
            try:
                employee = Employee.objects.get(user=user)
            except Employee.DoesNotExist:
                employee = None

            visibility_q = Q()
            if employee:
                visibility_q |= Q(role_assignments__employees=employee)
            if hasattr(Loan, "created_by"):
                visibility_q |= Q(created_by=user)

            assigner_exists = Task.objects.filter(loan_id=OuterRef("pk"), assigner_id=user.id)
            assignee_exists = Task.objects.filter(loan_id=OuterRef("pk"), assignee_id=employee.id) if employee else Task.objects.none()

            qs = qs.annotate(
                is_assigner=Exists(assigner_exists),
                is_assignee=Exists(assignee_exists)
            ).filter(
                visibility_q | Q(is_assigner=True) | Q(is_assignee=True)
            )

        # Exclude archived
        if self.request.query_params.get("include_archived", "").lower() != "true":
            qs = qs.filter(is_archived=False)

        # Filters
        broker_id = self.request.query_params.get("broker")
        officer_id = self.request.query_params.get("loan_officer")
        lead_id = self.request.query_params.get("team_leader")
        processor_id = self.request.query_params.get("processor")
        milestone_id = self.request.query_params.get("milestone")
        start_date = self.request.query_params.get("start_date")
        end_date = self.request.query_params.get("end_date")

        if broker_id:
            qs = qs.filter(broker_id=broker_id)
        if officer_id:
            qs = qs.filter(loan_officer_id=officer_id)
        if lead_id:
            qs = qs.filter(
                role_assignments__role__name__iexact='lead',
                role_assignments__employees__id=lead_id
            )

        if processor_id:
            qs = qs.filter(
                role_assignments__role__name__iexact='processor',
                role_assignments__employees__id=processor_id
            )

        if milestone_id:
            qs = qs.filter(milestone_id=milestone_id)
        if start_date:
            qs = qs.filter(created_at__date__gte=start_date)

        if end_date:
            qs = qs.filter(created_at__date__lte=end_date)
            
        # Prefetch role assignments for efficient access
        qs = qs.prefetch_related(
            Prefetch(
                'role_assignments',
                queryset=LoanRoleAssignment.objects.select_related('role').prefetch_related('employees')
            )
        )

        return qs.distinct()



    def list(self, request, *args, **kwargs):
        queryset = self.get_queryset()

        if not queryset.exists():
            return Response({
                "results": [],
                "message": "No data found for selected filters."
            }, status=200)

        milestone_counts = (
            queryset
            .filter(
                milestone__isnull=False,
                milestone__include_in_reports=True,   # ✅ USE EXISTING FIELD
                milestone__status="active",           # ✅ OPTIONAL but recommended
            )
            .values(
                "milestone__id",
                "milestone__name",
                "milestone__sort_order",
                "milestone__color",
                "milestone__background_color",
            )
            .annotate(count=Count("id"))
            .order_by("milestone__sort_order")
        )

        results = [
            {
                "milestone_id": row["milestone__id"],
                "milestone": row["milestone__name"],
                "count": row["count"],
            }
            for row in milestone_counts
        ]

        total_loans = sum(row["count"] for row in results)

        return Response({
            "results": results,
            
            "total_loans": total_loans,
        })


# -----------------------------
# LINKED EMPLOYEES BY BROKER
# -----------------------------
class BrokerLinkedEmployeesAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, *args, **kwargs):
        broker_id = request.query_params.get("broker")
        if not broker_id:
            return Response({
                "loan_officers": [],
                "team_leaders": [],
                "processors": []
            })

        loans = Loan.objects.filter(broker_id=broker_id).prefetch_related(
            Prefetch(
                "role_assignments",
                queryset=LoanRoleAssignment.objects.select_related("role").prefetch_related("employees")
            ),
            "loan_officer"
        )

        officers = {}
        team_leaders = {}
        processors = {}

        for loan in loans:
            if loan.loan_officer:
                officers[loan.loan_officer.id] = loan.loan_officer.name

            for ra in loan.role_assignments.all():
                if ra.role.name.lower() == "lead":
                    for emp in ra.employees.all():
                        team_leaders[emp.id] = emp.name

                if ra.role.name.lower() == "processor":
                    for emp in ra.employees.all():
                        processors[emp.id] = emp.name

        return Response({
            "loan_officers": [{"id": k, "name": v} for k, v in officers.items()],
            "team_leaders": [{"id": k, "name": v} for k, v in team_leaders.items()],
            "processors": [{"id": k, "name": v} for k, v in processors.items()],
        })


# -----------------------------
# LINKED PROCESSORS BY TEAM LEAD
# -----------------------------
class TeamLeadProcessorsAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, *args, **kwargs):
        team_lead_id = request.query_params.get("team_leader")
        if not team_lead_id:
            return Response({"processors": []})

        loans = Loan.objects.filter(
            role_assignments__role__name__icontains="lead",
            role_assignments__employees__id=team_lead_id
        ).prefetch_related(
            Prefetch(
                "role_assignments",
                queryset=LoanRoleAssignment.objects.select_related("role").prefetch_related("employees")
            )
        ).distinct()

        processors = {}

        for loan in loans:
            for ra in loan.role_assignments.all():
                if ra.role.name.lower() == "processor":
                    for emp in ra.employees.all():
                        processors[emp.id] = emp.name

        return Response({
            "processors": [{"id": k, "name": v} for k, v in processors.items()]
        })
    


class LoanExportReportAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def _auto_size_columns(self, ws, min_widths=None):
        min_widths = min_widths or {}

        # Header row is row 5
        headers = {cell.column: cell.value for cell in ws[5]}

        for col in ws.columns:
            max_length = 0
            col_letter = get_column_letter(col[0].column)
            header = headers.get(col[0].column, "")

            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))

            width = max_length + 3

            if header in min_widths:
                width = max(width, min_widths[header])

            ws.column_dimensions[col_letter].width = width

                
    def _get_period_label(self):
        if self.start_date and self.end_date:
            return f"{self.start_date.strftime('%b %Y')} – {self.end_date.strftime('%b %Y')}"
        if self.start_date:
            return f"From {self.start_date.strftime('%b %Y')}"
        if self.end_date:
            return f"Up to {self.end_date.strftime('%b %Y')}"
        return "All Time"


    def _add_invoice_header(self, ws, title, column_count):
        # Title
        ws.append([title])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=column_count)
        ws["A1"].font = Font(bold=True, size=14)

        # Generated On
        ws.append([f"Generated On: {datetime.now().strftime('%d %b %Y')}"])
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=column_count)

        # Period
        period_label = self._get_period_label()
        ws.append([f"Period: {period_label}"])
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=column_count)
        ws["A3"].font = Font(italic=True)

        # Spacer
        ws.append([])


    def _apply_date_filter(self, qs):
        if self.start_date:
            qs = qs.filter(created_at__date__gte=self.start_date)
        if self.end_date:
            qs = qs.filter(created_at__date__lte=self.end_date)
        return qs


    def _apply_visibility(self, qs, user):
        if user.is_superuser or user.has_perm("loan.view_all_loans"):
            return qs

        try:
            employee = Employee.objects.get(user=user)
        except Employee.DoesNotExist:
            return qs.none()

        visibility_q = Q(role_assignments__employees=employee)

        if hasattr(Loan, "created_by"):
            visibility_q |= Q(created_by=user)

        assigner_exists = Task.objects.filter(
            loan_id=OuterRef("pk"),
            assigner_id=user.id
        )

        assignee_exists = Task.objects.filter(
            loan_id=OuterRef("pk"),
            assignee_id=employee.id
        )

        return qs.annotate(
            is_assigner=Exists(assigner_exists),
            is_assignee=Exists(assignee_exists),
        ).filter(visibility_q | Q(is_assigner=True) | Q(is_assignee=True))
    
    def get(self, request):
        export_type = request.query_params.get("type")
        entity_id = request.query_params.get("id")
        start_date = request.query_params.get("start_date")
        end_date = request.query_params.get("end_date")

        self.start_date = (
            datetime.strptime(start_date, "%Y-%m-%d").date()
            if start_date else None
        )
        self.end_date = (
            datetime.strptime(end_date, "%Y-%m-%d").date()
            if end_date else None
        )

        if export_type == "broker":
            return self.export_broker_wise(entity_id)

        if export_type == "team_leader":
            return self.export_team_lead_wise(entity_id)

        if export_type == "processor":
            return self.export_processor_wise(entity_id)

        return self.export_all()

    

    def export_processor_wise(self, processor_id):
        wb = Workbook()
        ws = wb.active
        ws.title = "Processor Monthly"

        months = [
            "Jan", "Feb", "Mar", "Apr", "May", "Jun",
            "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"
        ]

        headers = ["Processor"] + months + ["Total"]

        self._add_invoice_header(
            ws,
            "Invoice – Processor Wise Monthly Report",
            len(headers)
        )

        ws.append(headers)
        ws.freeze_panes = "A4"

        processor = get_object_or_404(Employee, id=processor_id)

        base_qs = self._apply_visibility(
            Loan.objects.filter(
                role_assignments__role__name__iexact="processor",
                role_assignments__employees=processor
            ),
            self.request.user
        )
        base_qs = self._apply_date_filter(base_qs)

        monthly = (
            base_qs
            .filter(milestone__include_in_reports=True)
            .annotate(month=ExtractMonth("created_at"))
            .values("month")
            .annotate(total=Count("id"))
        )

        month_map = {i: 0 for i in range(1, 13)}
        for m in monthly:
            month_map[m["month"]] = m["total"]

        total = sum(month_map.values())

        row = [processor.name] + [month_map[i] for i in range(1, 13)] + [total]
        ws.append(row)

        return self.download(wb, "invoice_processor_monthly.xlsx")


    
    def _add_processor_sheet(self, wb):
        ws = wb.create_sheet("Processor Monthly")

        months = [
            "Jan", "Feb", "Mar", "Apr", "May", "Jun",
            "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"
        ]

        headers = ["Processor"] + months + ["Total", "Average"]

        self._add_invoice_header(
            ws,
            "Invoice – Processor Wise Monthly Summary",
            len(headers)
        )

        ws.append(headers)
        ws.freeze_panes = "A4"

        processors = Employee.objects.filter(
            roles__name__iexact="processor"
        ).distinct().order_by("name")

        base_qs = self._apply_visibility(
            Loan.objects.filter(
                role_assignments__role__name__iexact="processor"
            ),
            self.request.user
        )

        base_qs = self._apply_date_filter(base_qs)

        for processor in processors:
            qs = base_qs.filter(role_assignments__employees=processor)

            monthly = (
                qs
                .filter(milestone__include_in_reports=True)
                .annotate(month=ExtractMonth("created_at"))
                .values("month")
                .annotate(total=Count("id"))
            )

            month_map = {i: 0 for i in range(1, 13)}
            for m in monthly:
                month_map[m["month"]] = m["total"]

            total = sum(month_map.values())
            average = round(total / 12, 2)
            row = [processor.name] + [month_map[i] for i in range(1, 13)] + [total, average]
            ws.append(row)
        self._auto_size_columns(
            ws,
            min_widths={
                "Rate": 12,
                "Notes": 30,
                "State": 10,
            }
        )

    def _add_team_lead_sheet(self, wb):
        ws = wb.create_sheet("Team Lead Wise")

        months = [
            "Jan", "Feb", "Mar", "Apr", "May", "Jun",
            "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"
        ]

        headers = ["Team Lead"] + months + ["Total", "Average"]

        self._add_invoice_header(
            ws,
            "Invoice – Team Lead Wise Monthly Summary",
            len(headers)
        )

        ws.append(headers)
        ws.freeze_panes = "A4"

        team_leads = Employee.objects.filter(
            roles__name__iexact="lead"
        ).distinct().order_by("name")

        base_qs = self._apply_visibility(
            Loan.objects.filter(milestone__include_in_reports=True),
            self.request.user
        )

        base_qs = self._apply_date_filter(base_qs)

        for lead in team_leads:
            qs = base_qs.filter(
                role_assignments__role__name__iexact="lead",
                role_assignments__employees=lead
            )

            monthly = (
                qs.annotate(month=ExtractMonth("created_at"))
                .values("month")
                .annotate(total=Count("id"))
            )

            month_map = {i: 0 for i in range(1, 13)}
            for m in monthly:
                if m["month"]:
                    month_map[m["month"]] = m["total"]

            total = sum(month_map.values())
            average = round(total / 12, 2)
            row = [lead.name] + [month_map[i] for i in range(1, 13)] + [total, average]
            ws.append(row)

        
        self._auto_size_columns(
            ws,
            min_widths={
                "Rate": 12,
                "Notes": 30,
                "State": 10,
            }
        )



    def _add_broker_sheet(self, wb, broker_id=None):
        ws = wb.create_sheet("Broker Wise Loans")

        milestones = (
            Milestone.objects
            .filter(include_in_reports=True)
            .order_by("sort_order")
        )

        headers = (
            ["Broker", "Borrower", "Property", "Lender", "LO"]
            + [m.name for m in milestones]
            + ["Rate", "Notes", "State"]
        )

        self._add_invoice_header(
            ws,
            "Invoice – Broker Wise Loan Report",
            len(headers)
        )

        ws.append(headers)
        ws.freeze_panes = "A4"

        loans = self._apply_visibility(
            Loan.objects
                .select_related("broker", "loan_officer")
                .prefetch_related(
                    "lenders",
                    "milestone_history"  # ✅ correct relation
                ),
            self.request.user
        )

        loans = self._apply_date_filter(loans)

        if broker_id:
            loans = loans.filter(broker_id=broker_id)

        loans = loans.order_by("broker__name", "created_at")

        for loan in loans:
            # 🔹 Build milestone → latest date map
            milestone_date_map = {}

            for h in loan.milestone_history.all():
                if (
                    h.milestone_id not in milestone_date_map
                    or h.changed_at > milestone_date_map[h.milestone_id]
                ):
                    milestone_date_map[h.milestone_id] = h.changed_at

            # 🔹 Dates per milestone column
            milestone_dates = [
                milestone_date_map[m.id].strftime("%m/%d/%Y")
                if m.id in milestone_date_map
                else ""
                for m in milestones
            ]

            ws.append([
                loan.broker.name if loan.broker else "",
                f"{loan.first_name} {loan.last_name}",
                loan.subject_property or "",
                ", ".join(l.lender_name for l in loan.lenders.all()),
                loan.loan_officer.name if loan.loan_officer else "",
                *milestone_dates,
                "",  # Rate
                "",  # Notes
                "",  # State
            ])
            
        self._auto_size_columns(
                ws,
                min_widths={
                    "Rate": 12,
                    "Notes": 30,
                    "State": 10,
                }
            )

    def download(self, wb, filename):
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        wb.save(response)
        return response
    
    def export_all(self):
        wb = Workbook()
        wb.remove(wb.active)

        self._add_processor_sheet(wb)
        self._add_team_lead_sheet(wb)
        self._add_broker_sheet(wb)

        return self.download(wb, "invoice_all_loan_reports.xlsx")


    def export_team_lead_wise(self, team_lead_id=None):
        wb = Workbook()
        ws = wb.active
        ws.title = "Team Lead Wise"

        milestones = Milestone.objects.filter(
            include_in_reports=True
        ).order_by("sort_order")

        headers = ["Team Lead"] + [m.name for m in milestones] + ["Total"]

        self._add_invoice_header(
            ws,
            "Invoice – Team Lead Wise Loan Report",
            len(headers)
        )
        ws.append(headers)
        ws.freeze_panes = "A4"


        leads = Employee.objects.filter(
            roles__name__iexact="lead"
        ).distinct()

        if team_lead_id:
            leads = leads.filter(id=team_lead_id)

        for lead in leads:
            row = [lead.name]
            total = 0

            for m in milestones:
                qs = self._apply_visibility(
                    Loan.objects.filter(
                        milestone=m,
                        role_assignments__role__name__iexact="lead",
                        role_assignments__employees=lead
                    ),
                    self.request.user
                )

                qs = self._apply_date_filter(qs)

                count = qs.distinct().count()

                row.append(count)
                total += count

            row.append(total)
            ws.append(row)

        return self.download(wb, "team_lead_report.xlsx")


    def export_broker_wise(self, broker_id=None):
        wb = Workbook()
        wb.remove(wb.active)

        self._add_broker_sheet(wb, broker_id)

        return self.download(wb, "broker_wise_loans.xlsx")
